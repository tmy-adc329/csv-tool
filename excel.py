"""
excel.py — aggregate.py が作った4つの表を、1つのExcelに書き出す。
シート4枚、金額は3桁区切り、見出しは太字、列幅は内容に合わせる。
openpyxl に依存するのはこのファイルだけ。
"""

import os

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import config


def _display_width(text):
    """全角文字を2、半角文字を1として、表示上のだいたいの幅を返す。"""
    return sum(2 if ord(ch) > 255 else 1 for ch in str(text))


def _write_sheet(writer, sheet_name, frame):
    """DataFrame を1シートに書き、見出し太字・数値3桁区切り・列幅調整をする。"""
    frame.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]

    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

    for i, column_name in enumerate(frame.columns, start=1):
        widths = [_display_width(column_name)]
        for value in frame.iloc[:, i - 1]:
            widths.append(_display_width(value))
        ws.column_dimensions[get_column_letter(i)].width = max(widths) + 2


def write_report(by_shop, by_category, by_product, summary_frame, year_month):
    """4つの表を output/売上集計_YYYYMM.xlsx に書く。書いたパスを返す。"""
    os.makedirs(config.OUTPUT_DIR, exist_ok=True)
    path = os.path.join(config.OUTPUT_DIR, f"売上集計_{year_month}.xlsx")

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _write_sheet(writer, "店舗別", by_shop)
        _write_sheet(writer, "カテゴリ別", by_category)
        _write_sheet(writer, "商品別", by_product)
        _write_sheet(writer, "サマリ", summary_frame)

    return path
